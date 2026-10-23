"""Inspect an import source before binding columns in a Loginom package.

The MCP builder writes field names from whatever the caller passes. When the
caller guesses (`Summ` instead of `Summa_pokupki`, or a Cyrillic header instead
of the transliterated Loginom Name), the first downstream node fails with
«ошибка входного порта». This module reads the source file the way Loginom will
and reports the exact Name / DisplayName / DataType / DataKind to bind.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

from .loginom_ident import loginom_ident

TEXT_SUFFIXES = {".txt", ".csv", ".tsv", ".dat", ".tab"}
EXCEL_SUFFIXES = {".xlsx", ".xlsm"}

_DATE_FORMATS = (
    "%d.%m.%Y",
    "%d.%m.%y",
    "%Y-%m-%d",
    "%d/%m/%Y",
    "%m/%d/%Y",
    "%m/%d/%y",
    "%d.%m.%Y %H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%dT%H:%M:%S",
)
_DATE_RE = re.compile(
    r"^(\d{1,2}[./]\d{1,2}[./]\d{2,4}|\d{4}-\d{2}-\d{2})"
    r"([ T]\d{1,2}:\d{2}(:\d{2})?)?$"
)
_NULL_TOKENS = {"", "?", "null", "NULL", "None", "N/A", "n/a", "-"}
_TRUE_TOKENS = {"true", "false", "истина", "ложь", "да", "нет"}
_INT_RE = re.compile(r"^[+-]?\d+$")
_FLOAT_RE = re.compile(r"^[+-]?(\d+([.,]\d*)?|[.,]\d+)([eE][+-]?\d+)?$")

# Loginom defaults: float is continuous, everything else discrete.
_KIND_BY_TYPE = {
    "dtString": "dkDiscrete",
    "dtBoolean": "dkDiscrete",
    "dtDateTime": "dkDiscrete",
    "dtInteger": "dkDiscrete",
    "dtFloat": "dkContinuous",
}


def _decode(raw: bytes, encoding: str | None) -> tuple[str, str]:
    if raw.startswith(b"\xef\xbb\xbf"):
        return raw[3:].decode("utf-8", errors="replace"), "utf-8 (BOM)"
    if encoding:
        return raw.decode(encoding, errors="replace"), encoding
    try:
        return raw.decode("utf-8"), "utf-8"
    except UnicodeDecodeError:
        return raw.decode("cp1251", errors="replace"), "cp1251"


def _guess_delimiter(header: str) -> str:
    counts = {",": header.count(","), ";": header.count(";"), "\t": header.count("\t")}
    best = max(counts, key=counts.get)
    return best if counts[best] else "\t"


def _read_head(p: Path, *, sample_rows: int, max_bytes: int = 8 * 1024 * 1024) -> bytes:
    """Read only the head of a large file: header plus ``sample_rows`` lines."""
    need = sample_rows + 1
    chunks: list[bytes] = []
    lines = 0
    total = 0
    with p.open("rb") as fh:
        while total < max_bytes and lines < need:
            chunk = fh.read(65536)
            if not chunk:
                break
            chunks.append(chunk)
            total += len(chunk)
            lines += chunk.count(b"\n")
    raw = b"".join(chunks)
    if lines >= need:
        cut = raw.rfind(b"\n")
        if cut >= 0:
            raw = raw[: cut + 1]
    return raw


def read_text_preview(
    path: str | Path,
    *,
    encoding: str | None = None,
    delimiter: str | None = None,
    sample_rows: int = 500,
) -> dict:
    """Read header + a sample of rows. Supports .txt/.csv with any delimiter."""
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(str(p))
    if p.suffix.lower() not in TEXT_SUFFIXES:
        raise ValueError(
            f"{p.suffix or 'file'} is not a delimited text source; "
            "supported: " + ", ".join(sorted(TEXT_SUFFIXES | EXCEL_SUFFIXES))
        )
    text, used_encoding = _decode(
        _read_head(p, sample_rows=sample_rows), encoding
    )
    lines = text.splitlines()[: sample_rows + 1]
    if not lines:
        raise ValueError("file is empty")
    delim = delimiter or _guess_delimiter(lines[0])
    reader = csv.reader(lines, delimiter=delim, quotechar='"')
    rows = list(reader)
    header = [h.strip().strip('"') for h in rows[0]]
    body = rows[1:]
    width = len(header)
    body = [r + [""] * (width - len(r)) if len(r) < width else r[:width] for r in body]
    return {
        "path": str(p.resolve()),
        "encoding": used_encoding,
        "delimiter": delim,
        "header": header,
        "rows": body,
        "rows_sampled": len(body),
    }


def read_excel_preview(
    path: str | Path,
    *,
    sheet: int | str | None = 1,
    sample_rows: int = 200,
    max_columns: int = 200,
) -> dict:
    """Header + sample rows of an .xlsx/.xlsm sheet via openpyxl (read-only)."""
    p = Path(path)
    if not p.is_file():
        raise FileNotFoundError(str(p))
    if p.suffix.lower() not in EXCEL_SUFFIXES:
        raise ValueError(
            f"{p.suffix or 'file'} is not an Excel source; "
            "supported: " + ", ".join(sorted(EXCEL_SUFFIXES))
        )
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on install
        raise RuntimeError(
            "openpyxl is required to inspect .xlsx sources; "
            "install it or pass columns=[...] explicitly"
        ) from exc
    wb = load_workbook(filename=str(p), read_only=True, data_only=True)
    try:
        if isinstance(sheet, str):
            ws = wb[sheet]
        else:
            ws = wb.worksheets[max(int(sheet or 1) - 1, 0)]
        it = ws.iter_rows(values_only=True)
        try:
            raw_header = next(it)
        except StopIteration:
            raise ValueError("sheet is empty")
        header = [
            (str(v).strip() if v is not None else f"Column{i + 1}")
            for i, v in enumerate(raw_header)
        ][:max_columns]
        body: list[list[str]] = []
        for row in it:
            body.append(
                ["" if v is None else str(v) for v in row[: len(header)]]
            )
            if len(body) >= sample_rows:
                break
    finally:
        wb.close()
    width = len(header)
    body = [r + [""] * (width - len(r)) for r in body]
    return {
        "path": str(p.resolve()),
        "encoding": None,
        "delimiter": None,
        "header": header,
        "rows": body,
        "rows_sampled": len(body),
        "sheet": ws.title,
        "sheet_number": wb.sheetnames.index(ws.title) + 1,
    }


def _is_null(value: str) -> bool:
    return value.strip().strip('"') in _NULL_TOKENS


def infer_column_type(values: list[str]) -> tuple[str, str]:
    """Guess (DataType, DataKind) from sample values. Nulls are ignored."""
    clean = [v.strip().strip('"') for v in values if not _is_null(v)]
    if not clean:
        return "dtString", "dkDiscrete"
    if all(_INT_RE.match(v) for v in clean):
        return "dtInteger", _KIND_BY_TYPE["dtInteger"]
    if all(_FLOAT_RE.match(v) for v in clean):
        return "dtFloat", _KIND_BY_TYPE["dtFloat"]
    if all(_DATE_RE.match(v) for v in clean):
        return "dtDateTime", _KIND_BY_TYPE["dtDateTime"]
    if all(v.lower() in _TRUE_TOKENS for v in clean):
        return "dtBoolean", _KIND_BY_TYPE["dtBoolean"]
    return "dtString", _KIND_BY_TYPE["dtString"]


def _sample_values(values: list[str], limit: int = 3) -> list[str]:
    seen: list[str] = []
    for v in values:
        v = v.strip()
        if v and v not in seen:
            seen.append(v[:60])
        if len(seen) >= limit:
            break
    return seen


def inspect_source_schema(
    path: str | Path,
    *,
    encoding: str | None = None,
    delimiter: str | None = None,
    sample_rows: int = 500,
    max_columns: int = 200,
    sheet: int | str | None = 1,
) -> dict:
    """Field list as Loginom will expose it: Name, DisplayName, DataType, DataKind."""
    if Path(path).suffix.lower() in EXCEL_SUFFIXES:
        preview = read_excel_preview(
            path, sheet=sheet, sample_rows=min(sample_rows, 200), max_columns=max_columns
        )
    else:
        preview = read_text_preview(
            path, encoding=encoding, delimiter=delimiter, sample_rows=sample_rows
        )
    header = preview["header"]
    if len(header) > max_columns:
        raise ValueError(f"too many columns in header: {len(header)}")
    columns = []
    for i, label in enumerate(header):
        values = [row[i] for row in preview["rows"]]
        data_type, data_kind = infer_column_type(values)
        columns.append(
            {
                "name": loginom_ident(label),
                "display_name": label,
                "data_type": data_type,
                "data_kind": data_kind,
                "nulls_in_sample": sum(1 for v in values if _is_null(v)),
                "sample": _sample_values(values),
            }
        )
    result = {
        "ok": True,
        "path": preview["path"],
        "encoding": preview["encoding"],
        "delimiter": preview["delimiter"],
        "column_count": len(columns),
        "rows_sampled": preview["rows_sampled"],
        "columns": columns,
        "note": (
            "Types here are inferred from a sample; override with "
            "configure_node(columns=[...]). configure_node(source_file=...) copies "
            "code page and delimiter into import_settings for text sources, and "
            "sheet_number for Excel."
        ),
    }
    if preview.get("sheet"):
        result["sheet"] = preview["sheet"]
    if preview.get("sheet_number") is not None:
        result["sheet_number"] = preview["sheet_number"]
    return result


def columns_from_schema(
    schema: dict, overrides: list[dict] | None = None
) -> list[dict]:
    """Import-ready ColumnDefs ({name, display_name, data_type, data_kind})."""
    by_label: dict[str, dict] = {}
    for col in schema["columns"]:
        by_label[col["display_name"]] = col
        by_label[col["name"]] = col
    out: list[dict] = []
    for col in schema["columns"]:
        out.append(
            {
                "name": col["name"],
                "display_name": col["display_name"],
                "data_type": col["data_type"],
                "data_kind": col["data_kind"],
            }
        )
    for override in overrides or []:
        label = override.get("name")
        base = by_label.get(label)
        if base is None and label:
            # Allow an explicit column that is not in the header sample.
            out.append(dict(override))
            continue
        if base is not None:
            for item in out:
                if item["name"] in (base["display_name"], base["name"]):
                    for key in ("data_type", "data_kind", "display_name"):
                        if override.get(key):
                            item[key] = override[key]
                    break
    return out


def columns_for_import(
    path: str | Path,
    *,
    encoding: str | None = None,
    delimiter: str | None = None,
    sample_rows: int = 500,
    overrides: list[dict] | None = None,
    sheet: int | str | None = 1,
) -> list[dict]:
    """Import-ready ColumnDefs ({name, display_name, data_type, data_kind})."""
    schema = inspect_source_schema(
        path,
        encoding=encoding,
        delimiter=delimiter,
        sample_rows=sample_rows,
        sheet=sheet,
    )
    return columns_from_schema(schema, overrides)
